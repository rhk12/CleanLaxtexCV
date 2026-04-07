from __future__ import annotations

import re
import subprocess
from pathlib import Path
from typing import Optional

import unidecode
from docx import Document


class DossierError(Exception):
    pass


def paragraphs_text(file_path: str | Path) -> str:
    doc = Document(str(file_path))
    full_text = [unidecode.unidecode(para.text) for para in doc.paragraphs]
    return "\n".join(full_text)


def get_table_data(doc: Document, table_index: int) -> list[dict[str, str]]:
    table_data: list[dict[str, str]] = []
    table = doc.tables[table_index]
    keys: Optional[tuple[str, ...]] = None

    for row in table.rows:
        text = tuple(cell.text for cell in row.cells)
        is_bold_row = all(
            all(run.bold for run in cell.paragraphs[0].runs if run.text.strip())
            for cell in row.cells
            if cell.paragraphs
        )

        if is_bold_row:
            keys = text
            continue

        if keys:
            table_data.append(dict(zip(keys, text)))

    return table_data


def extract_text_between_markers(full_text: str, start_marker: str, end_marker: str | None = None) -> str:
    start_marker = re.escape(start_marker)
    if end_marker:
        end_marker = re.escape(end_marker)
        pattern = re.compile(f"{start_marker}(.*?){end_marker}", re.DOTALL)
    else:
        pattern = re.compile(f"{start_marker}(.*)", re.DOTALL)

    match = pattern.search(full_text)
    return match.group(1).strip() if match else ""


def ensure_docx(input_path: str | Path, verbose: bool = False) -> Path:
    path = Path(input_path)
    if not path.exists():
        raise DossierError(f"Dossier not found: {path}")

    if path.suffix.lower() == ".docx":
        return path

    if path.suffix.lower() != ".doc":
        raise DossierError(f"Unsupported dossier extension: {path.suffix}. Expected .doc or .docx")

    converted = path.with_suffix(".docx")
    if converted.exists() and converted.stat().st_mtime >= path.stat().st_mtime:
        if verbose:
            print(f"[info] Using existing converted docx: {converted}")
        return converted

    soffice = shutil_which("soffice")
    if not soffice:
        raise DossierError(
            "A .doc dossier was provided, but python-docx only reads .docx. "
            "Install LibreOffice (soffice) or save the dossier as .docx first."
        )

    cmd = [soffice, "--headless", "--convert-to", "docx", str(path), "--outdir", str(path.parent)]
    if verbose:
        print("[info] Converting .doc to .docx")
        print("[debug]", " ".join(cmd))
    subprocess.run(cmd, check=True)

    if not converted.exists():
        raise DossierError(f"Conversion completed but no .docx was created for {path}")

    return converted


def find_latest_dossier(dossiers_dir: str | Path, verbose: bool = False) -> Path:
    directory = Path(dossiers_dir)
    if not directory.exists():
        raise DossierError(f"Dossiers directory not found: {directory}")

    candidates = sorted(
        list(directory.glob("*.docx")) + list(directory.glob("*.doc")),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise DossierError(f"No .doc or .docx files found in {directory}")

    if verbose:
        print(f"[info] Selected latest dossier: {candidates[0]}")
    return candidates[0]


def shutil_which(cmd: str) -> Optional[str]:
    from shutil import which
    return which(cmd)
